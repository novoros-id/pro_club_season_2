"""Explicit live smoke test for the MWS-backed audio-to-DOCX pipeline."""

import argparse
import json
import math
import sys
import time
from pathlib import Path


PREP_DIR = Path(__file__).resolve().parents[1]
if str(PREP_DIR) not in sys.path:
    sys.path.insert(0, str(PREP_DIR))

from docx import Document  # noqa: E402

from create_file.create_docx import create_docx, table_segments_time  # noqa: E402
from model_api import MWSClient  # noqa: E402
from text_to_paragraphs.text_to_paragraphs import cosine_similarity  # noqa: E402
from transcription_audio.transcription import Transcription  # noqa: E402


EMBEDDING_INPUTS = [
    "Пользователь открывает справочник номенклатуры.",
    "Затем пользователь выбирает необходимый элемент.",
    "Сегодня на улице идёт сильный дождь.",
]


class CountingMWSClient(MWSClient):
    """MWS client with non-sensitive per-operation counters for smoke assertions."""

    def __init__(self):
        super().__init__()
        self.calls = {"chat": 0, "embeddings": 0, "transcription": 0}

    def chat_completion(self, *args, **kwargs):
        self.calls["chat"] += 1
        return super().chat_completion(*args, **kwargs)

    def embed_texts(self, *args, **kwargs):
        self.calls["embeddings"] += 1
        return super().embed_texts(*args, **kwargs)

    def transcribe(self, *args, **kwargs):
        self.calls["transcription"] += 1
        return super().transcribe(*args, **kwargs)


def run_smoke(audio_path, output_dir):
    audio = Path(audio_path).expanduser().resolve()
    output = Path(output_dir).expanduser().resolve()
    if not audio.is_file():
        raise FileNotFoundError(f"Audio file not found: {audio}")
    output.mkdir(parents=True, exist_ok=True)

    client = CountingMWSClient()
    if not client.config.api_key:
        raise RuntimeError("MWS_API_KEY is not configured.")
    print("MWS_API_KEY configured: yes")

    started = time.perf_counter()
    reply = client.chat_completion(
        [{"role": "user", "content": "Ответь только одним словом: OK"}],
        model="gemma-4-31b-it",
    )
    print(
        "CHAT "
        f"model=gemma-4-31b-it status={client.last_response_status} "
        f"duration={time.perf_counter() - started:.2f}s preview={reply[:12]!r}"
    )

    started = time.perf_counter()
    vectors = client.embed_texts(EMBEDDING_INPUTS)
    dimensions = {len(vector) for vector in vectors}
    finite = all(math.isfinite(value) for vector in vectors for value in vector)
    if len(vectors) != len(EMBEDDING_INPUTS) or len(dimensions) != 1 or not finite:
        raise RuntimeError("Embedding contract validation failed.")
    print(
        "EMBED "
        f"model=bge-m3 status={client.last_response_status} "
        f"duration={time.perf_counter() - started:.2f}s "
        f"count={len(vectors)} dimension={next(iter(dimensions))}"
    )
    print(
        "EMBED similarities "
        f"1-2={cosine_similarity(vectors[0], vectors[1]):.6f} "
        f"2-3={cosine_similarity(vectors[1], vectors[2]):.6f}"
    )

    json_path = output / f"{audio.stem}.whisper.json"
    started = time.perf_counter()
    transcriber = Transcription(client=client)
    transcriber.save_json(str(audio), str(json_path))
    transcription = json.loads(json_path.read_text(encoding="utf-8"))
    segments = transcription.get("segments", [])
    if not transcription.get("full_text") or not segments:
        raise RuntimeError("Transcription does not contain full_text and segments.")
    print(
        "STT "
        f"model=whisper-large-v3 status={client.last_response_status} "
        f"duration={time.perf_counter() - started:.2f}s "
        f"text_length={len(transcription['full_text'])} segments={len(segments)} "
        f"fields={sorted(segments[0])} "
        f"first={segments[0]['start']}-{segments[0]['end']} "
        f"last={segments[-1]['start']}-{segments[-1]['end']}"
    )

    sentence_rows = table_segments_time(str(json_path))
    if not sentence_rows:
        raise RuntimeError("No paragraphs can be formed from transcription segments.")
    calls_before_docx = dict(client.calls)
    started = time.perf_counter()
    docx_path = Path(create_docx(str(json_path), "", True, client=client).get_docx())
    document = Document(docx_path)
    nonempty = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip()
        and paragraph.style
        and paragraph.style.name.startswith("Heading")
    ]
    if not docx_path.is_file() or not nonempty or not headings:
        raise RuntimeError("Generated DOCX did not pass content validation.")
    if client.calls["embeddings"] <= calls_before_docx["embeddings"]:
        raise RuntimeError("DOCX pipeline did not call BGE-M3.")
    if client.calls["chat"] <= calls_before_docx["chat"]:
        raise RuntimeError("DOCX pipeline did not call Gemma.")
    print(
        "DOCX "
        f"path={docx_path} paragraphs={len(nonempty)} headings={len(headings)} "
        f"duration={time.perf_counter() - started:.2f}s"
    )
    print(
        "CALLS "
        f"chat={client.calls['chat']} embeddings={client.calls['embeddings']} "
        f"transcription={client.calls['transcription']}"
    )
    return docx_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audio", required=True, help="Path to an MP3/WAV/FLAC/OGG file")
    parser.add_argument("--output-dir", required=True, help="Directory for generated JSON and DOCX")
    args = parser.parse_args()
    run_smoke(args.audio, args.output_dir)


if __name__ == "__main__":
    main()
